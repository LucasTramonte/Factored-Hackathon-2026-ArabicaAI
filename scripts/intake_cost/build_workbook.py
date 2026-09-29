"""Build the editable, source-linked intake capacity and cost workbook behind ADR-004.

Every derived cell is a live formula over the ``Inputs`` sheet, so reviewers can change an
assumption and see the effect. ``python scripts/intake_cost/build_workbook.py --print`` also
prints the same figures computed in Python; ADR-004 quotes those, so the ADR and the workbook
can't drift.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "Docs/Costs/Intake/INTAKE_COST_ESTIMATE.xlsx"
CHECKED = "2026-09-29"
NAVY, TEAL, AMBER, WHITE = "17343B", "147D76", "FFF0C2", "FFFFFF"

# (key, label, value, unit, source or definition, evidence status)
INPUTS = [
    ("s1", "S1 volume: p95 daily 'Cargo no reconocido' complaints (2025)", 17, "episodes/day", "Silver fact_complaints, 2025 daily p95", "observed synthetic"),
    ("s2", "S2 volume: p95 daily 'Queja' interactions", 145, "episodes/day", "Silver fact_call_center_interactions, daily p95", "observed synthetic"),
    ("s3", "S3 volume: p95 daily call-center interactions (all reasons)", 818, "episodes/day", "Silver fact_call_center_interactions, daily p95 (max 894)", "observed synthetic"),
    ("s4", "S4 volume: 10x S3 stress", 8180, "episodes/day", "Stress test only; production volume is not disclosed", "assumption"),
    ("window_days", "Evaluation window", 32, "days", "2026-09-29 to 2026-10-31", "fact"),
    ("month_days", "Days per modeled month", 30, "days", "Comparable 30-day month", "assumption"),
    ("peak", "Busiest-hour factor over the daily average", 3, "x", "Synthetic hourly profile is flat (~4.2%/hour); 3x is a stress assumption", "assumption"),
    ("doc_req", "Worker requests per page load (gated HTML document)", 1, "requests/episode", "run_worker_first covers /, /index.html, /agent", "code-derived"),
    ("cust_req", "Customer API requests per episode", 3, "requests/episode", "login, list charges, create case (budget test)", "measured locally"),
    ("cust_q", "D1 queries per customer episode", 9, "queries/episode", "back-end/test/integration/budget.test.js", "measured locally"),
    ("cust_read", "D1 rows read per customer episode", 10, "rows/episode", "budget test, local D1", "measured locally"),
    ("cust_write", "D1 rows written per customer episode (incl. index writes)", 7, "rows/episode", "budget test, local D1", "measured locally"),
    ("agent_refresh", "Agent queue refreshes per accepted case", 1, "refreshes/case", "One agent look per case", "assumption"),
    ("agent_req", "Worker requests per agent refresh", 2, "requests/refresh", "agent session + case list", "code-derived"),
    ("agent_read", "D1 rows read per agent refresh (full 50-case page)", 155, "rows/refresh", "Upper bound: 3 x 51 joined rows + session; 13 measured with few cases", "derived bound"),
    ("agent_write", "D1 rows written per agent refresh", 3, "rows/refresh", "budget test, local D1", "measured locally"),
    ("bytes_case", "Stored bytes per case, typical 77-character statement", 367, "bytes", "SQLite with the Worker migrations, 10,000 cases, after VACUUM", "measured locally"),
    ("bytes_case_max", "Stored bytes per case, 2,000-character statement", 4268, "bytes", "Same method, worst case", "measured locally"),
    ("cpu_ms", "Worker CPU per request", 5, "ms", "Not measured; read the real value from Workers analytics after deploy", "assumption"),
    ("cf_req_day", "Workers Free: requests per day", 100000, "requests/day", "developers.cloudflare.com/workers/platform/pricing", "official limit"),
    ("cf_cpu_req", "Workers Free: CPU per request", 10, "ms", "same", "official limit"),
    ("d1_read_day", "D1 Free: rows read per day", 5000000, "rows/day", "same", "official limit"),
    ("d1_write_day", "D1 Free: rows written per day", 100000, "rows/day", "same", "official limit"),
    ("d1_db_mb", "D1 Free: max database size", 500, "MB", "developers.cloudflare.com/d1/platform/limits", "official limit"),
    ("d1_q_inv", "D1 Free: queries per Worker invocation", 50, "queries", "same", "official limit"),
    ("trigger", "Upgrade trigger: share of any daily Free limit", 0.7, "share", "ADR-004 decision; sustained 3 days", "decision"),
    ("paid_base", "Workers Paid base", 5, "USD/month", "Workers pricing", "official price"),
    ("paid_req_incl", "Paid: included requests", 10000000, "requests/month", "same", "official price"),
    ("paid_req_rate", "Paid: extra requests", 0.30, "USD/million", "same", "official price"),
    ("paid_cpu_incl", "Paid: included CPU", 30000000, "ms/month", "same", "official price"),
    ("paid_cpu_rate", "Paid: extra CPU", 0.02, "USD/million ms", "same", "official price"),
    ("paid_read_incl", "Paid: included D1 rows read", 25000000000, "rows/month", "same", "official price"),
    ("paid_read_rate", "Paid: extra D1 rows read", 0.001, "USD/million", "same", "official price"),
    ("paid_write_incl", "Paid: included D1 rows written", 50000000, "rows/month", "same", "official price"),
    ("paid_write_rate", "Paid: extra D1 rows written", 1.00, "USD/million", "same", "official price"),
    ("lambda_req_rate", "AWS Lambda requests", 0.20, "USD/million", "aws.amazon.com/lambda/pricing", "official price"),
    ("lambda_gbs_rate", "AWS Lambda compute (x86)", 0.0000166667, "USD/GB-s", "same", "official price"),
    ("lambda_free_req", "Lambda always-free requests", 1000000, "requests/month", "same", "official price"),
    ("lambda_free_gbs", "Lambda always-free compute", 400000, "GB-s/month", "same", "official price"),
    ("lambda_mem", "Lambda memory", 0.5, "GB", "Sizing assumption", "assumption"),
    ("lambda_dur", "Lambda duration per request", 0.1, "s", "Not measured", "assumption"),
    ("apigw_rate", "API Gateway HTTP API", 1.00, "USD/million", "aws.amazon.com/api-gateway/pricing (1M/month free for 12 months, not applied)", "official price"),
    ("ddb_wru_rate", "DynamoDB on-demand writes", 0.625, "USD/million WRU", "aws.amazon.com/dynamodb/pricing/on-demand", "official price"),
    ("ddb_rru_rate", "DynamoDB on-demand reads", 0.125, "USD/million RRU", "same", "official price"),
    ("ddb_read_ratio", "RRU per D1 row read (eventually consistent, <4 KB)", 0.5, "RRU/row", "Mapping assumption from the D1 row model", "assumption"),
    ("nat_hour", "NAT gateway", 0.045, "USD/hour", "aws.amazon.com/vpc/pricing", "official price"),
    ("ipv4_hour", "Public IPv4 address", 0.005, "USD/hour", "same", "official price"),
    ("rds_month", "RDS PostgreSQL db.t4g.micro, Single-AZ, on-demand", 11.68, "USD/month", "Indicative: price not shown on the fetched page; confirm in the AWS Pricing Calculator", "indicative"),
    ("rds_storage", "RDS 20 GB general-purpose storage", 2.30, "USD/month", "Indicative; confirm in the AWS Pricing Calculator", "indicative"),
    ("hours_month", "Hours per month", 730, "hours", "AWS billing convention", "fact"),
    ("ai_in", "Future AI input tokens per episode", 12000, "tokens", "Not measured; no model runs in the MVP (ADR-002)", "assumption"),
    ("ai_out", "Future AI output tokens per episode", 2000, "tokens", "Not measured", "assumption"),
    ("wai_free", "Workers AI free allocation", 10000, "neurons/day", "developers.cloudflare.com/workers-ai/platform/pricing", "official price"),
    ("wai_rate", "Workers AI beyond the allocation", 0.011, "USD/1,000 neurons", "same", "official price"),
    ("oss_in", "gpt-oss-20b input", 18182, "neurons/M tokens", "same", "official price"),
    ("oss_out", "gpt-oss-20b output", 27273, "neurons/M tokens", "same", "official price"),
    ("llama_in", "llama-3.3-70b-instruct-fp8-fast input", 26668, "neurons/M tokens", "same", "official price"),
    ("llama_out", "llama-3.3-70b-instruct-fp8-fast output", 204805, "neurons/M tokens", "same", "official price"),
    ("haiku_in", "Claude Haiku 4.5 input", 1.00, "USD/M tokens", "Anthropic API list price", "official price"),
    ("haiku_out", "Claude Haiku 4.5 output", 5.00, "USD/M tokens", "same", "official price"),
    ("sonnet_in", "Claude Sonnet 5 input", 2.00, "USD/M tokens", "same", "official price"),
    ("sonnet_out", "Claude Sonnet 5 output", 10.00, "USD/M tokens", "same", "official price"),
]
FIRST_INPUT_ROW = 4
REF = {key: f"Inputs!$B${FIRST_INPUT_ROW + i}" for i, (key, *_rest) in enumerate(INPUTS)}
VALUE = {key: value for key, _label, value, *_rest in INPUTS}
SCENARIOS = ("s1", "s2", "s3", "s4")

SOURCES = [
    ("Cloudflare Workers pricing and limits", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ("Cloudflare D1 limits", "https://developers.cloudflare.com/d1/platform/limits/"),
    ("Cloudflare Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/"),
    ("Cloudflare Zero Trust plans (Access free to 50 users)", "https://www.cloudflare.com/plans/zero-trust-services/"),
    ("AWS Free plan", "https://docs.aws.amazon.com/awsaccountbilling/latest/aboutv2/free-tier-plans.html"),
    ("AWS Lambda pricing", "https://aws.amazon.com/lambda/pricing/"),
    ("Amazon API Gateway pricing", "https://aws.amazon.com/api-gateway/pricing/"),
    ("Amazon DynamoDB on-demand pricing", "https://aws.amazon.com/dynamodb/pricing/on-demand/"),
    ("Amazon VPC pricing (NAT, IPv4)", "https://aws.amazon.com/vpc/pricing/"),
    ("Amazon CloudFront pricing (Free flat-rate plan)", "https://aws.amazon.com/cloudfront/pricing/"),
    ("Amazon RDS for PostgreSQL pricing", "https://aws.amazon.com/rds/postgresql/pricing/"),
    ("AWS Pricing Calculator", "https://calculator.aws/"),
]


def model(v: dict) -> dict:
    """The same arithmetic as the workbook formulas, for printing and for ADR-004."""
    out = {}
    per_req = v["doc_req"] + v["cust_req"] + v["agent_req"] * v["agent_refresh"]
    per_read = v["cust_read"] + v["agent_read"] * v["agent_refresh"]
    per_write = v["cust_write"] + v["agent_write"] * v["agent_refresh"]
    api_req = v["cust_req"] + v["agent_req"] * v["agent_refresh"]
    fixed_o4 = v["rds_month"] + v["rds_storage"] + (v["nat_hour"] + v["ipv4_hour"]) * v["hours_month"]
    for s in SCENARIOS:
        d = v[s]
        req, read, write = d * per_req, d * per_read, d * per_write
        util = {"requests": req / v["cf_req_day"], "rows_read": read / v["d1_read_day"], "rows_written": write / v["d1_write_day"]}
        m = v["month_days"]
        paid = (v["paid_base"] + max(0, req * m - v["paid_req_incl"]) * v["paid_req_rate"] / 1e6
                + max(0, req * m * v["cpu_ms"] - v["paid_cpu_incl"]) * v["paid_cpu_rate"] / 1e6
                + max(0, read * m - v["paid_read_incl"]) * v["paid_read_rate"] / 1e6
                + max(0, write * m - v["paid_write_incl"]) * v["paid_write_rate"] / 1e6)
        lam = d * api_req * m
        aws = (max(0, lam - v["lambda_free_req"]) * v["lambda_req_rate"] / 1e6
               + max(0, lam * v["lambda_mem"] * v["lambda_dur"] - v["lambda_free_gbs"]) * v["lambda_gbs_rate"]
               + lam * v["apigw_rate"] / 1e6
               + write * m * v["ddb_wru_rate"] / 1e6 + read * m * v["ddb_read_ratio"] * v["ddb_rru_rate"] / 1e6)
        out[s] = {
            "episodes_day": d, "worker_req_day": req, "rows_read_day": read, "rows_written_day": write,
            "util": util, "binding": max(util, key=util.get), "max_util": max(util.values()),
            "free_ok": max(util.values()) <= v["trigger"],
            "storage_mb_window": d * v["window_days"] * v["bytes_case"] / 1e6,
            "storage_mb_window_max": d * v["window_days"] * v["bytes_case_max"] / 1e6,
            "peak_writes_s": write / 86400 * v["peak"],
            "cf_paid_month": paid, "aws_serverless_month": aws, "aws_o4_month": aws + fixed_o4,
            "ai_month": {name: d * m * cost for name, cost in ai_per_episode(v).items()},
        }
    out["capacity_free_episodes_day"] = min(v["cf_req_day"] / per_req, v["d1_read_day"] / per_read, v["d1_write_day"] / per_write)
    out["per_episode"] = {"worker_requests": per_req, "rows_read": per_read, "rows_written": per_write}
    out["o4_fixed_month"] = fixed_o4
    out["ai_per_episode"] = ai_per_episode(v)
    out["ai_free_episodes_day"] = {k: v["wai_free"] / n for k, n in neurons_per_episode(v).items()}
    return out


def neurons_per_episode(v: dict) -> dict:
    return {"workers_ai_gpt_oss_20b": (v["ai_in"] * v["oss_in"] + v["ai_out"] * v["oss_out"]) / 1e6,
            "workers_ai_llama_3_3_70b": (v["ai_in"] * v["llama_in"] + v["ai_out"] * v["llama_out"]) / 1e6}


def ai_per_episode(v: dict) -> dict:
    n = neurons_per_episode(v)
    return {"workers_ai_gpt_oss_20b": n["workers_ai_gpt_oss_20b"] * v["wai_rate"] / 1000,
            "workers_ai_llama_3_3_70b": n["workers_ai_llama_3_3_70b"] * v["wai_rate"] / 1000,
            "claude_haiku_4_5": (v["ai_in"] * v["haiku_in"] + v["ai_out"] * v["haiku_out"]) / 1e6,
            "claude_sonnet_5": (v["ai_in"] * v["sonnet_in"] + v["ai_out"] * v["sonnet_out"]) / 1e6}


def title(ws, text: str, last_col: int, subtitle: str | None = None) -> None:
    """Title row, optional italic subtitle, no gridlines."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(size=16, bold=True, color=WHITE)
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.cell(2, 1, subtitle).font = Font(italic=True, color=NAVY)
    ws.sheet_view.showGridLines = False


def header(ws, row: int, labels) -> None:
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row, col, label)
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.font = Font(bold=True, color=WHITE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30


def grid(ws, first_row: int, rows: list[tuple], widths: dict, formats: dict | None = None) -> None:
    for r, values in enumerate(rows, start=first_row):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, value)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if formats and c in formats:
                cell.number_format = formats[c]
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build() -> Workbook:
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    r = REF

    readme = wb.active
    readme.title = "Read me"
    title(readme, "ArabicaAI | Intake capacity and cost (ADR-004)", 2)
    notes = [
        ("Decision", "Run the intake service on Cloudflare Workers + D1 Free through 2026-10-31. Move to Workers Paid ($5/month) if any daily Free limit stays above 70% for 3 days."),
        ("Window", "2026-09-29 to 2026-10-31. Finalists are announced on 2026-10-15."),
        ("Demand", "Observed synthetic sample only: interactions p50 664, p95 818, max 894 per day. Production volume is not disclosed, so S1-S4 are sizing scenarios, not forecasts."),
        ("Measured", "D1 queries, rows and stored bytes come from local tests (back-end/test/integration/budget.test.js; SQLite with the Worker migrations). Worker CPU is not measured yet."),
        ("Editing", "Change the yellow cells on Inputs. Every other sheet recalculates."),
        ("Not included", "Taxes, FX, custom domain, paid logs, egress beyond the free allowances, labour, support, high availability, security operations, credits."),
        ("AWS figures", "Per-request AWS prices come from official pages. RDS rows are indicative and must be confirmed in the AWS Pricing Calculator (procedure in ADR-004)."),
        ("AI", "No model runs in the MVP (ADR-002). The AI sheet is a future envelope with unmeasured tokens, not a plan."),
        ("Checked", f"Prices and limits checked {CHECKED}."),
    ]
    grid(readme, 3, notes, {"A": 16, "B": 120})

    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs | change yellow cells only with evidence", 5)
    header(ws, 3, ("Parameter", "Value", "Unit", "Source or definition", "Evidence"))
    for i, (_key, label, value, unit, source, status) in enumerate(INPUTS):
        row = FIRST_INPUT_ROW + i
        for col, v in enumerate((label, value, unit, source, status), start=1):
            ws.cell(row, col, v).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=AMBER)
        ws.cell(row, 2).font = Font(bold=True, color=NAVY)
    for col, width in {"A": 52, "B": 16, "C": 18, "D": 70, "E": 18}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B4"

    per_req = f"({r['doc_req']}+{r['cust_req']}+{r['agent_req']}*{r['agent_refresh']})"
    per_read = f"({r['cust_read']}+{r['agent_read']}*{r['agent_refresh']})"
    per_write = f"({r['cust_write']}+{r['agent_write']}*{r['agent_refresh']})"
    api_req = f"({r['cust_req']}+{r['agent_req']}*{r['agent_refresh']})"

    cf = wb.create_sheet("Cloudflare capacity")
    title(cf, "Workers + D1 Free: daily usage per scenario and binding constraint", 12,
          "A customer episode plus one agent refresh per case. Utilization above the trigger means Workers Paid.")
    header(cf, 4, ("Scenario", "Episodes/day", "Worker requests/day", "% of 100k requests", "D1 rows read/day", "% of 5M reads",
                   "D1 rows written/day", "% of 100k writes", "Highest utilization", "Within Free (<= trigger)?",
                   "Cases stored by 2026-10-31 (MB, typical)", "Peak D1 writes/s"))
    for i, s in enumerate(SCENARIOS):
        row = 5 + i
        cells = [s.upper(), f"={r[s]}", f"=B{row}*{per_req}", f"=C{row}/{r['cf_req_day']}", f"=B{row}*{per_read}",
                 f"=E{row}/{r['d1_read_day']}", f"=B{row}*{per_write}", f"=G{row}/{r['d1_write_day']}",
                 f"=MAX(D{row},F{row},H{row})", f'=IF(I{row}<={r["trigger"]},"yes","no: upgrade")',
                 f"=B{row}*{r['window_days']}*{r['bytes_case']}/1000000", f"=G{row}/86400*{r['peak']}"]
        grid(cf, row, [tuple(cells)], {}, {2: "#,##0", 3: "#,##0", 4: "0.0%", 5: "#,##0", 6: "0.0%", 7: "#,##0",
                                            8: "0.0%", 9: "0.0%", 11: "#,##0.0", 12: "0.00"})
    cf.cell(10, 1, "Free capacity (episodes/day)").font = Font(bold=True)
    cf.cell(10, 2, f"=MIN({r['cf_req_day']}/{per_req},{r['d1_read_day']}/{per_read},{r['d1_write_day']}/{per_write})").number_format = "#,##0"
    for col, width in zip("ABCDEFGHIJKL", (12, 13, 15, 12, 15, 12, 15, 12, 13, 16, 18, 12)):
        cf.column_dimensions[col].width = width

    cost = wb.create_sheet("Monthly cost")
    title(cost, "Monthly cost per scenario (USD, before tax)", 6,
          "Cloudflare Paid includes its $5 base. The AWS serverless equivalent is API Gateway HTTP + Lambda + DynamoDB on-demand, with CloudFront Free for static files.")
    header(cost, 4, ("Scenario", "Cloudflare Free", "Cloudflare Workers Paid", "AWS serverless equivalent",
                     "AWS O4: Lambda + RDS + NAT (indicative)", "Note"))
    m = r["month_days"]
    for i, s in enumerate(SCENARIOS):
        row = 5 + i
        cap = f"'Cloudflare capacity'!"
        req, read, write = f"{cap}C{row}", f"{cap}E{row}", f"{cap}G{row}"
        paid = (f"={r['paid_base']}+MAX(0,{req}*{m}-{r['paid_req_incl']})*{r['paid_req_rate']}/1000000"
                f"+MAX(0,{req}*{m}*{r['cpu_ms']}-{r['paid_cpu_incl']})*{r['paid_cpu_rate']}/1000000"
                f"+MAX(0,{read}*{m}-{r['paid_read_incl']})*{r['paid_read_rate']}/1000000"
                f"+MAX(0,{write}*{m}-{r['paid_write_incl']})*{r['paid_write_rate']}/1000000")
        lam = f"({r[s]}*{api_req}*{m})"
        aws = (f"=MAX(0,{lam}-{r['lambda_free_req']})*{r['lambda_req_rate']}/1000000"
               f"+MAX(0,{lam}*{r['lambda_mem']}*{r['lambda_dur']}-{r['lambda_free_gbs']})*{r['lambda_gbs_rate']}"
               f"+{lam}*{r['apigw_rate']}/1000000+{write}*{m}*{r['ddb_wru_rate']}/1000000"
               f"+{read}*{m}*{r['ddb_read_ratio']}*{r['ddb_rru_rate']}/1000000")
        o4 = f"=D{row}+{r['rds_month']}+{r['rds_storage']}+({r['nat_hour']}+{r['ipv4_hour']})*{r['hours_month']}"
        free = f'=IF(\'Cloudflare capacity\'!I{row}<={r["trigger"]},0,"exceeds Free")'
        grid(cost, row, [(s.upper(), free, paid, aws, o4, "RDS rows are indicative; confirm in the AWS Pricing Calculator")], {},
             {2: '"$"#,##0.00', 3: '"$"#,##0.00', 4: '"$"#,##0.00', 5: '"$"#,##0.00'})
    for col, width in zip("ABCDEF", (12, 16, 20, 22, 28, 58)):
        cost.column_dimensions[col].width = width

    ai = wb.create_sheet("AI options (future)")
    title(ai, "Future AI envelope: not in the MVP; tokens are unmeasured", 7)
    header(ai, 4, ("Model", "USD per episode", "Free episodes/day", "S1 USD/month", "S2 USD/month", "S3 USD/month", "S4 USD/month"))
    per = {
        "Workers AI gpt-oss-20b": (f"(({r['ai_in']}*{r['oss_in']}+{r['ai_out']}*{r['oss_out']})/1000000)", True),
        "Workers AI llama-3.3-70b": (f"(({r['ai_in']}*{r['llama_in']}+{r['ai_out']}*{r['llama_out']})/1000000)", True),
        "Claude Haiku 4.5": (f"(({r['ai_in']}*{r['haiku_in']}+{r['ai_out']}*{r['haiku_out']})/1000000)", False),
        "Claude Sonnet 5": (f"(({r['ai_in']}*{r['sonnet_in']}+{r['ai_out']}*{r['sonnet_out']})/1000000)", False),
    }
    for i, (name, (expr, neurons)) in enumerate(per.items()):
        row = 5 + i
        usd = f"{expr}*{r['wai_rate']}/1000" if neurons else expr
        free = f"={r['wai_free']}/{expr}" if neurons else "none"
        grid(ai, row, [(name, f"={usd}", free, *[f"={usd}*{r[s]}*{m}" for s in SCENARIOS])], {},
             {2: '"$"#,##0.0000', 3: "#,##0", 4: '"$"#,##0.00', 5: '"$"#,##0.00', 6: '"$"#,##0.00', 7: '"$"#,##0.00'})
    ai.cell(10, 1, "Claude via Amazon Bedrock: take the rate from the AWS Pricing Calculator; whether credits apply is unverified.")
    for col, width in zip("ABCDEFG", (28, 16, 18, 14, 14, 14, 14)):
        ai.column_dimensions[col].width = width

    src = wb.create_sheet("Sources")
    title(src, f"Sources (checked {CHECKED})", 2)
    header(src, 3, ("Source", "URL"))
    grid(src, 4, SOURCES, {"A": 50, "B": 90})
    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--print", action="store_true", help="print the figures ADR-004 quotes")
    args = parser.parse_args()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")
    if args.print:
        import json
        print(json.dumps(model(VALUE), indent=2, default=float))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
